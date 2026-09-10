"""Currency invariants from invoice to refund, credit note, ledger and exports."""
from datetime import date
from decimal import Decimal
import os
from pathlib import Path
import unittest
from unittest.mock import patch

from desktop.tests import test_critical_workflows as fixtures


class CurrencyWorkflowTests(unittest.TestCase):
    setUp = fixtures.CriticalWorkflowTests.setUp
    tearDown = fixtures.CriticalWorkflowTests.tearDown
    invoice_payload = fixtures.CriticalWorkflowTests.invoice_payload
    standard_items = staticmethod(fixtures.CriticalWorkflowTests.standard_items)

    def issue(self, currency):
        invoice_id = self.db.save_invoice(self.invoice_payload(currency=currency), self.standard_items())
        self.db.save_invoice(self.invoice_payload(id=invoice_id, currency=currency, status_code="pending_approval"), self.standard_items())
        self.db.approve_invoice(invoice_id, "QA Owner")
        self.db.save_invoice(self.invoice_payload(id=invoice_id, currency=currency, status_code="issued", prepared_by_role="owner"), self.standard_items())
        return invoice_id

    def credit(self, currency):
        invoice_id = self.issue(currency)
        self.db.add_payment(invoice_id, date.today(), 1200, "Banka")
        self.db.add_payment_refund(invoice_id, date.today(), 240, "Banka", "QA refund")
        note_id = self.db.create_credit_note(invoice_id, date.today(), 240, "QA credit")
        return invoice_id, self.db.credit_note_export_payload(note_id)

    def test_credit_refund_and_correction_keep_source_currency(self):
        for currency in ("RSD", "EUR", "BGN"):
            with self.subTest(currency=currency):
                invoice_id, note = self.credit(currency)
                self.assertEqual(note["currency"], currency)
                self.assertEqual(note["source_invoice"]["currency"], currency)
                self.assertEqual(note["net_amount"], 200)
                self.assertEqual(note["vat_amount"], 40)
                self.assertEqual(self.db.credit_note_draft_info(invoice_id)["available_gross"], 0)
                self.assertEqual(self.db.prepare_invoice_correction_draft(invoice_id)["currency"], currency)
                with self.assertRaises(ValueError):
                    self.db.create_credit_note(invoice_id, date.today(), 1, "Duplicate refund")

    def test_reports_use_company_currency_and_never_add_other_currencies(self):
        company = self.db.get_company()
        self.db.save_company({**company, "country_code": "RS", "default_currency": "RSD"})
        for currency in ("RSD", "EUR"):
            self.credit(currency)
            self.db.save_project_document({"project_id": self.project_id, "document_type": "input", "document_date": date.today(), "partner_name": "QA supplier", "net_amount": 100, "vat_rate": .2, "currency": currency})
        for currency in (None, "RSD", "EUR"):
            report = self.db.project_accountant_report(self.project_id, date.today(), date.today(), currency=currency)
            self.assertEqual(report["currency"], currency or "RSD")
            self.assertEqual(report["totals"]["output_net"], 800)
            self.assertEqual(report["totals"]["input_net"], 100)
            self.assertEqual(report["totals"]["vat_payable"], 140)
            self.assertEqual(report["totals"]["net_collected"], 960)
            self.assertEqual(len(report["foreign_currency_rows"]), 3)
            self.assertEqual(len(report["foreign_currency_payments"]), 2)
            summary = self.db.project_period_summary(self.project_id, date.today(), date.today(), currency=currency)
            self.assertEqual(summary["currency"], currency or "RSD")
            self.assertEqual(summary["profit_net"], 700)
            self.assertEqual(summary["excluded_currency_count"], 5)
        with self.assertRaises(ValueError):
            self.db.project_vat_evidence(self.project_id, date.today(), date.today(), currency="INVALID")

    def test_credit_note_cannot_bypass_write_or_period_controls(self):
        invoice_id = self.issue("RSD")
        self.db.add_payment(invoice_id, date.today(), 1200, "Banka")
        self.db.add_payment_refund(invoice_id, date.today(), 240, "Banka", "QA refund")
        for guard in ("assert_business_write_access", "assert_financial_date_open"):
            with self.subTest(guard=guard), patch.object(self.db, guard, side_effect=ValueError("Blocked by control")):
                with self.assertRaisesRegex(ValueError, "Blocked by control"):
                    self.db.create_credit_note(invoice_id, date.today(), 240, "QA blocked")
                self.assertEqual(self.db.list_credit_notes(invoice_id=invoice_id), [])

    def test_credit_is_capped_by_invoice_even_after_an_overpayment(self):
        invoice_id = self.issue("RSD")
        self.db.add_payment(invoice_id, date.today(), 2400, "Banka")
        self.db.add_payment_refund(invoice_id, date.today(), 2400, "Banka", "QA excess payment returned")
        self.assertEqual(self.db.credit_note_draft_info(invoice_id)["available_gross"], 1200)
        with self.assertRaises(ValueError):
            self.db.create_credit_note(invoice_id, date.today(), 2400, "More than original invoice")
        self.db.create_credit_note(invoice_id, date.today(), 1200, "Full original invoice")
        self.assertEqual(self.db.credit_note_draft_info(invoice_id)["available_gross"], 0)

    def test_rsd_exports_and_mixed_currency_cancellations(self):
        import delta_fakture_export as exporter
        from openpyxl import load_workbook
        from pypdf import PdfReader

        _, note = self.credit("RSD")
        report = self.db.project_accountant_report(self.project_id, date.today(), date.today(), currency="RSD")
        report["report_language"] = "bg"
        report["company"]["name"] = "Тестова фирма - дугачак назив предузећа"
        report["cancelled_rows"] = [{"document_no": "QA-CANCEL-EUR", "currency": "EUR", "gross_amount": 999, "net_amount": 999, "vat_amount": 0}]
        output = Path(os.environ.get("OPSNEST_QA_OUTPUT") or self.temp_dir.name)
        output.mkdir(parents=True, exist_ok=True)
        note_xlsx = exporter.export_credit_note_xlsx(note, output / "credit-RSD.xlsx")
        book = load_workbook(note_xlsx)
        formats = [c.number_format for s in book for row in s for c in row if 'RSD' in c.number_format]
        self.assertTrue(formats)
        self.assertFalse(any('EUR' in c.number_format for s in book for row in s for c in row))
        book.close()
        report_xlsx = exporter.export_project_accountant_xlsx(report, output / "accountant-RSD.xlsx")
        book = load_workbook(report_xlsx)
        corrections = book[exporter.report_text("sheet_corrections", "bg")]
        euro_row = next(row for row in corrections if any(c.value == "QA-CANCEL-EUR" for c in row))
        self.assertIn('EUR', euro_row[-1].number_format)
        summary = book[exporter.report_text("sheet_summary", "bg")]
        self.assertIn('RSD', summary["B11"].number_format)
        book.close()
        for name, factory, payload in (
            ("credit-RSD", exporter.export_credit_note_pdf, note),
            ("vat-RSD", exporter.export_project_vat_evidence_pdf, report),
            ("accountant-RSD", exporter.export_project_accountant_pdf, report),
        ):
            pdf = factory(payload, output / (name + ".pdf"))
            pages = PdfReader(pdf).pages
            self.assertTrue(pages)
            self.assertIn("RSD", " ".join(page.extract_text() for page in pages))
            if name != "credit-RSD":
                content = " ".join(" ".join(page.extract_text() for page in pages).split())
                self.assertNotIn("Uz fakturu", content)
                self.assertIn("Към фактура", content)
                self.assertIn("QA credit", content)
                if name == "accountant-RSD":
                    self.assertNotIn("Povraćaj -", content)
                    self.assertIn("Банков превод", content)
            for page in pages:
                self.assertAlmostEqual(float(page.mediabox.width), 595.276, delta=1)
                self.assertAlmostEqual(float(page.mediabox.height), 841.89, delta=1)


if __name__ == "__main__":
    unittest.main()
